"""Parser for the client's monthly car-wash workbook.

The workbook has one sheet per day of the month, named "1".."31", plus a few
month-level sheets we ignore. Each daily sheet is several blocks laid out side by
side; see CLAUDE.md for the full map. This module turns one workbook into a list
of `ParsedDay`s, each carrying a fingerprint so repeated uploads of the same
month can tell which days actually changed.

Nothing here touches the database — see `routers/imports.py` for that.
"""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field
from datetime import date as date_type
from decimal import Decimal, InvalidOperation

import openpyxl

from app.models import CASH_DENOMINATIONS

# Column letters, resolved to 1-based indexes below.
COL = {
    "sl": 1,          # A
    "reg_no": 2,      # B
    "mobile": 3,      # C
    "category": 4,    # D  -> car model
    "service": 5,     # E
    "wash_cash": 7,   # G
    "wash_gpay": 8,   # H
    "wash_credit": 9, # I
    "paint_cash": 13,   # M
    "paint_gpay": 14,   # N
    "paint_credit": 15, # O
    "puc_label": 19,  # S
    "puc_value": 20,  # T
    "exp_label": 23,  # W
    "exp_cash": 24,   # X
    "exp_gpay": 25,   # Y
    "sum_label": 27,  # AA
    "sum_value": 28,  # AB
    "denom": 30,      # AD
    "denom_qty": 31,  # AE
}

MAX_ROW = 80
MAX_COL = 32

# Rows in the register whose "reg no" looks like CRD30/6/26 are credit repayments,
# not service jobs. We don't model repayment, so they're skipped and reported.
CREDIT_ROW_PREFIX = "CRD"

# Summary labels whose values together make up the day's cash takings.
CASH_IN_LABELS = (
    "WASH (CASH)",
    "WASH CRDT RECVD (CASH)",
    "PAINT (CASH)",
    "PAINT CRDT RECVD (CASH)",
    "PUC (CASH)",
    "PUC CRDT RECVD (CASH)",
)


class ExcelImportError(Exception):
    """Raised when the uploaded file isn't the workbook we expect."""


@dataclass
class ParsedDay:
    date: date_type
    car_entries: list[dict] = field(default_factory=list)
    expenses: list[dict] = field(default_factory=list)
    puc: dict | None = None
    day_closing: dict | None = None
    skipped_credit_rows: int = 0

    @property
    def is_empty(self) -> bool:
        return not (self.car_entries or self.expenses or self.puc or self.day_closing)

    def fingerprint(self) -> str:
        """Stable hash of everything we'd write for this day.

        Lets a re-upload skip days that haven't changed, and rewrite the ones that
        have, without comparing row by row.
        """
        payload = {
            "date": self.date.isoformat(),
            "car_entries": self.car_entries,
            "expenses": self.expenses,
            "puc": self.puc,
            "day_closing": self.day_closing,
        }
        blob = json.dumps(payload, sort_keys=True, default=str, separators=(",", ":"))
        return hashlib.sha256(blob.encode()).hexdigest()


def _num(value) -> Decimal:
    """Coerce a cell to a Decimal, treating blanks and stray text as zero."""
    if value is None or value == "":
        return Decimal("0")
    if isinstance(value, bool):
        return Decimal("0")
    try:
        return Decimal(str(value).strip().replace(",", ""))
    except (InvalidOperation, ValueError):
        return Decimal("0")


def _text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _label(value) -> str:
    """Normalise a label cell so lookups survive stray spacing and case."""
    return " ".join(_text(value).upper().split())


def _grid(sheet) -> dict[int, dict[int, object]]:
    """Read a sheet into {row: {col: value}}, skipping blanks."""
    grid: dict[int, dict[int, object]] = {}
    for r, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=MAX_ROW, max_col=MAX_COL, values_only=True), start=1
    ):
        cells = {c: v for c, v in enumerate(row, start=1) if v is not None and _text(v) != ""}
        if cells:
            grid[r] = cells
    return grid


def _find_label_rows(grid, col: int, wanted: str) -> list[int]:
    """Every row where `col` holds `wanted`, top to bottom."""
    return [r for r in sorted(grid) if _label(grid[r].get(col)) == wanted]


def _parse_register(grid) -> tuple[list[dict], int]:
    """The car wash / paint register — one row per vehicle.

    Paint and wash are the same thing as far as this app is concerned, so their
    money columns are simply added together.
    """
    entries: list[dict] = []
    skipped = 0
    for r in sorted(grid):
        if r < 3:
            continue
        cells = grid[r]
        sl = cells.get(COL["sl"])
        # Register rows are the ones numbered in column A; the totals row isn't.
        if not isinstance(sl, (int, float)) or isinstance(sl, bool):
            continue

        reg_no = _text(cells.get(COL["reg_no"]))
        paid = (
            _num(cells.get(COL["wash_cash"]))
            + _num(cells.get(COL["wash_gpay"]))
            + _num(cells.get(COL["paint_cash"]))
            + _num(cells.get(COL["paint_gpay"]))
        )
        pending = _num(cells.get(COL["wash_credit"])) + _num(cells.get(COL["paint_credit"]))

        if not reg_no and paid == 0 and pending == 0:
            continue
        if reg_no.upper().startswith(CREDIT_ROW_PREFIX):
            # A repayment against an older credit, not a job done today.
            skipped += 1
            continue

        entries.append(
            {
                "reg_no": reg_no,
                "phone": _text(cells.get(COL["mobile"])),
                "car_model": _text(cells.get(COL["category"])),
                "service_type": _text(cells.get(COL["service"])),
                "amount_paid": str(paid),
                "amount_pending": str(pending),
            }
        )
    return entries, skipped


def _parse_puc(grid) -> dict | None:
    """PUC is recorded as daily totals, never per vehicle."""
    values = {}
    for wanted in ("CASH", "G PAY", "DISCOUNT"):
        rows = _find_label_rows(grid, COL["puc_label"], wanted)
        values[wanted] = _num(grid[rows[0]].get(COL["puc_value"])) if rows else Decimal("0")

    collection = values["CASH"] + values["G PAY"]
    discount = values["DISCOUNT"]
    if collection == 0 and discount == 0:
        return None
    return {"collection_amount": str(collection), "discount": str(discount)}


def _expense_total_rows(grid) -> list[int]:
    return _find_label_rows(grid, COL["exp_label"], "TOTAL")


def _parse_expenses(grid) -> list[dict]:
    """Expense lines run from row 3 down to the block's TOTAL row.

    Salary advances sit in their own block below that total; they're deliberately
    out of scope, so we stop at the first total.
    """
    totals = _expense_total_rows(grid)
    stop = totals[0] if totals else MAX_ROW

    expenses = []
    for r in sorted(grid):
        if r < 3 or r >= stop:
            continue
        cells = grid[r]
        description = _text(cells.get(COL["exp_label"]))
        amount = _num(cells.get(COL["exp_cash"])) + _num(cells.get(COL["exp_gpay"]))
        if not description or amount == 0:
            continue
        expenses.append({"description": description, "amount": str(amount)})
    return expenses


def _parse_denominations(grid) -> dict[str, int]:
    allowed = {str(d) for d in CASH_DENOMINATIONS}
    denominations: dict[str, int] = {}
    for r in sorted(grid):
        cells = grid[r]
        denom = cells.get(COL["denom"])
        if not isinstance(denom, (int, float)) or isinstance(denom, bool):
            continue
        key = str(int(denom))
        if key not in allowed:
            continue
        qty = int(_num(cells.get(COL["denom_qty"])))
        if qty:
            denominations[key] = qty
    return denominations


def _parse_day_closing(grid) -> dict | None:
    """The end-of-day cash and bank block."""
    label_col, value_col = COL["sum_label"], COL["sum_value"]

    def summary(wanted: str, occurrence: int = 0) -> Decimal:
        rows = _find_label_rows(grid, label_col, wanted)
        if len(rows) <= occurrence:
            return Decimal("0")
        return _num(grid[rows[occurrence]].get(value_col))

    cash_collected = sum((summary(lbl) for lbl in CASH_IN_LABELS), Decimal("0"))

    # Cash paid out: the expense block's cash total, plus the salary advance
    # block's cash total underneath it.
    totals = _expense_total_rows(grid)
    cash_expenses = sum(
        (_num(grid[r].get(COL["exp_cash"])) for r in totals), Decimal("0")
    )

    # "YESTERDAY CLOSING" / "TODAY CLOSING" appear twice: cash box first, bank second.
    opening_cash = summary("YESTERDAY CLOSING", 0)
    opening_bank = summary("YESTERDAY CLOSING", 1)
    closing_bank = summary("TODAY CLOSING", 1)
    cash_to_partner = summary("TO GEORGE")
    denominations = _parse_denominations(grid)

    if not denominations and opening_cash == 0 and cash_collected == 0 and closing_bank == 0:
        return None

    return {
        "opening_cash": str(opening_cash),
        "cash_collected": str(cash_collected),
        "cash_expenses": str(cash_expenses),
        "cash_to_partner": str(cash_to_partner),
        "opening_bank": str(opening_bank),
        "closing_bank": str(closing_bank),
        "denominations": denominations,
    }


def parse_workbook(path_or_stream, year: int, month: int) -> list[ParsedDay]:
    """Parse every filled-in day sheet for the given month.

    The month comes from the user rather than the file, because the sheets are
    named only "1".."31". Days that haven't been filled in yet are left out
    entirely, so uploading a half-finished month imports only what exists.
    """
    try:
        workbook = openpyxl.load_workbook(
            path_or_stream, data_only=True, read_only=True
        )
    except Exception as exc:  # openpyxl raises a grab-bag of errors on bad input
        raise ExcelImportError(f"Could not read that file as an Excel workbook: {exc}") from exc

    try:
        day_sheets = [n for n in workbook.sheetnames if n.strip().isdigit()]
        if not day_sheets:
            raise ExcelImportError(
                "This workbook has no day sheets. Expected sheets named 1 to 31."
            )

        days: list[ParsedDay] = []
        for name in day_sheets:
            day_number = int(name.strip())
            try:
                day_date = date_type(year, month, day_number)
            except ValueError:
                # e.g. a "31" sheet in a 30-day month — not a real date, skip it.
                continue

            grid = _grid(workbook[name])
            car_entries, skipped = _parse_register(grid)
            parsed = ParsedDay(
                date=day_date,
                car_entries=car_entries,
                expenses=_parse_expenses(grid),
                puc=_parse_puc(grid),
                day_closing=_parse_day_closing(grid),
                skipped_credit_rows=skipped,
            )
            if not parsed.is_empty:
                days.append(parsed)

        return days
    finally:
        workbook.close()
